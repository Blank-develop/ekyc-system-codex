from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/chilanhouthnitvongkhay/Downloads/ekyc-system-codex")
OUT_DIR = ROOT / "outputs" / "project_proposal"
SCREEN_DIR = OUT_DIR / "screenshots"
DOCX_PATH = OUT_DIR / "LALIGENCE_eKYC_Project_Proposal.docx"
LOGO = ROOT / "logo.png"

NAVY = RGBColor(8, 22, 50)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(95, 111, 139)
GOLD = RGBColor(231, 191, 53)
LIGHT_BLUE = "F4F6F9"
LIGHT_GOLD = "FFF7DE"
WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=9360, indent_dxa=120) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_para(doc, text="", size=11, color=NAVY, bold=False, italic=False, after=8, align=None, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, color=BLUE, bold=True)
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(10)
    elif level == 2:
        set_run_font(run, size=13, color=BLUE, bold=True)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        set_run_font(run, size=12, color=DARK_BLUE, bold=True)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    run = p.add_run(text)
    set_run_font(run, size=11, color=NAVY)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    run = p.add_run(text)
    set_run_font(run, size=11, color=NAVY)
    return p


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GOLD)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, color=NAVY, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(text)
    set_run_font(r2, size=10.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table)
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.5, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(value)
            set_run_font(r, size=9.3, color=NAVY)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def add_screenshot(doc, title, image_name, caption):
    add_heading(doc, title, level=2)
    img_path = SCREEN_DIR / image_name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(img_path), width=Inches(6.2))
    p.paragraph_format.space_after = Pt(4)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption)
    set_run_font(r, size=9.5, color=MUTED, italic=True)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = NAVY
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("LALIGENCE eKYC Proposal")
    set_run_font(run, size=9, color=MUTED)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    p.add_run().add_picture(str(LOGO), width=Inches(1.45))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    r = title.add_run("LALIGENCE eKYC Project Proposal")
    set_run_font(r, size=24, color=NAVY, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run("Secure document verification, liveness, face matching, and face-based payment access")
    set_run_font(r, size=13, color=MUTED)

    add_table(
        doc,
        ["Prepared for", "Prepared by", "Date", "Status"],
        [["Customer / Internal Review", "LALIGENCE eKYC Team", date.today().isoformat(), "Proposal draft"]],
        widths=[1.7, 1.7, 1.2, 1.6],
    )

    add_callout(
        doc,
        "Proposal objective",
        "Build and validate an IAL2-aligned eKYC workflow that verifies identity documents, checks fraud signals, confirms live user presence, matches the selfie to the document portrait, and enables returning face login or Face Pay authorization after enrollment.",
    )

    add_heading(doc, "1. Executive Summary", 1)
    add_para(
        doc,
        "LALIGENCE is a web-first eKYC and secure access platform for identity verification and face-based authorization. The proposed system guides a user through document proofing, active liveness, hand gesture challenge, selfie capture, passive liveness, face matching, and final decisioning. After a successful verification, the verified profile can be bound to one active Face ID for future login or payment authorization.",
    )
    add_para(
        doc,
        "The app is designed for financial services, banking onboarding, internal KYC checks, customer identity proofing, and secure face-based access. The current target is NIST IAL2-aligned readiness rather than a claim of formal certification.",
    )

    add_heading(doc, "2. Project Goals", 1)
    for item in [
        "Provide a guided, user-friendly verification flow for passport and Lao ID card evidence.",
        "Detect common document fraud signals such as edited photo, edited expiry, visual/MRZ mismatch, non-document upload, and poor-quality capture.",
        "Use active liveness and hand gestures to reduce static photo and replay attacks.",
        "Use passive liveness and face matching to compare the live selfie with the document portrait.",
        "Store only verification results, verified profile attributes, and one active face template per user according to privacy policy.",
        "Support returning face login and Face Pay transfer authorization for verified users.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. Proposed Solution", 1)
    add_para(
        doc,
        "The proposed solution contains a React web application and a Python backend. The frontend manages the capture experience, step gating, status display, and customer-facing flows. The backend performs document analysis, OCR/MRZ extraction, fraud scoring, liveness analysis, face matching, profile storage, and face-login matching.",
    )
    add_table(
        doc,
        ["Layer", "Role", "Key functions"],
        [
            ["Frontend", "User experience and capture", "Document upload/camera, active liveness, gesture UI, selfie capture, face login, payment authorization"],
            ["Backend", "Verification intelligence", "Document OCR, fraud scoring, face matching, anti-spoof checks, profile and session APIs"],
            ["Database", "Verification and profile records", "One user ID to one active face ID, one passport number to one verified profile, audit-ready results"],
            ["Models", "Local/open-source AI", "Face detection/recognition, anti-spoofing, OCR/MRZ, document quality and fraud checks"],
        ],
        widths=[1.1, 1.8, 3.3],
    )

    add_heading(doc, "4. End-to-End Workflow", 1)
    workflow = [
        "New user starts signup and receives or enters a user ID.",
        "User uploads or captures passport/Lao ID evidence.",
        "System checks document quality, document likeness, OCR/MRZ, expiry, and fraud signals.",
        "If the document passes, user performs active face liveness challenges.",
        "User completes hand gesture challenges inside a moving target circle.",
        "User captures a selfie for passive liveness and face matching.",
        "System returns final decision, scores, reason codes, and OCR summary.",
        "If accepted, the user can enroll Face ID for future face login and Face Pay authorization.",
    ]
    for item in workflow:
        add_number(doc, item)

    add_heading(doc, "5. Usage Guide With Screenshots", 1)
    add_screenshot(doc, "Start Page", "01-start-page.png", "Landing screen with Face Pay transfer, new user signup, face login, and core capability cards.")
    add_screenshot(doc, "New User Signup", "02-new-user-signup.png", "New users start verification with a generated or entered user ID.")
    add_screenshot(doc, "Document Proofing", "03-document-proofing.png", "The user uploads or captures passport/Lao ID evidence. The decision panel displays document quality, fraud risk, and OCR results.")
    add_screenshot(doc, "Active Face Liveness", "04-active-liveness.png", "The user performs camera-based liveness actions such as blink, head movement, and mouth opening.")
    add_screenshot(doc, "Hand Gesture Challenge", "05-gesture-challenge.png", "The user performs requested hand gestures in a moving target area.")
    add_screenshot(doc, "Selfie Face Match", "06-selfie-face-match.png", "The user captures a selfie for passive liveness and face matching against the document portrait.")
    add_screenshot(doc, "Result Decision", "07-result-decision.png", "The final result shows the decision, OCR, biometric scores, reason codes, and Face ID enrollment action.")
    add_screenshot(doc, "Face Login", "08-face-login.png", "Returning verified users can use live face login instead of repeating full onboarding.")
    add_screenshot(doc, "Face Pay Transfer", "09-face-pay-transfer.png", "Payment-style authorization requires a verified face scan before transfer approval.")

    add_heading(doc, "6. Security and Fraud Detection Scope", 1)
    add_table(
        doc,
        ["Control area", "Examples covered", "Expected behavior"],
        [
            ["Document fraud", "Edited portrait, edited expiry, text overlay, fake template, non-document upload", "Reject or route to manual review with reason codes"],
            ["Document validity", "Expiry, MRZ checksum, OCR/MRZ consistency, document number", "Reject invalid, expired, or inconsistent evidence"],
            ["Active liveness", "Blink, mouth open, head turn, screen replay attempts", "Pass only when the correct live human action is detected"],
            ["Gesture challenge", "Random moving target with simple hand signs", "Pass only when the correct gesture is detected inside the target"],
            ["Passive liveness", "Paper photo, phone screen, laptop screen, multiple faces", "Reject spoof or replay risk"],
            ["Face matching", "Same person vs impostor selfie", "Accept above threshold; reject low-match or multiple-face cases"],
        ],
        widths=[1.5, 2.7, 2.0],
    )

    add_heading(doc, "7. NIST IAL2 Alignment", 1)
    add_para(
        doc,
        "The implementation is intended to support NIST SP 800-63A-4 IAL2-aligned identity proofing. The current application supports evidence collection, attribute extraction, evidence validation, biometric ownership verification, liveness checks, workflow gating, and audit-style reason codes.",
    )
    add_callout(
        doc,
        "Compliance note",
        "A formal IAL2 claim requires more than UI and models. The organization must also document source validation, privacy notices, retention, deletion, audit logging, notification of proofing, operating procedures, and security controls.",
    )

    add_heading(doc, "8. Implementation Plan", 1)
    add_table(
        doc,
        ["Phase", "Deliverables", "Success criteria"],
        [
            ["Phase 1: Web pilot", "Frontend and backend demo, document proofing, liveness, selfie, result page", "Colleagues can complete end-to-end tests on web"],
            ["Phase 2: Model tuning", "Dataset tests for document fraud, face match, passive liveness, active liveness", "Measured false accept/reject rates and tuned thresholds"],
            ["Phase 3: Profile and payment flow", "Face ID enrollment, face login, Face Pay transfer authorization", "One user has one active Face ID and payment scan works"],
            ["Phase 4: Compliance readiness", "NIST IAL2 test evidence, audit logs, privacy/retention docs, security controls", "Readiness package prepared for customer review"],
            ["Phase 5: Production hardening", "Cloud deployment, database, monitoring, rate limiting, backup, incident response", "Stable multi-user testing with acceptable latency"],
        ],
        widths=[1.5, 2.7, 2.0],
    )

    add_heading(doc, "9. Deliverables", 1)
    for item in [
        "Web eKYC application with document, liveness, gesture, selfie, result, face login, and Face Pay screens.",
        "Python backend APIs for verification sessions, document upload, active liveness, selfie analysis, face enrollment, and face login.",
        "PostgreSQL-backed profile storage with SQLite fallback for local development.",
        "NIST IAL2 readiness test case matrix and execution evidence template.",
        "Deployment guidance for internal network, Cloudflare Tunnel, Render, Vercel, and future DigitalOcean/VPS production hosting.",
        "Flutter SDK package for mobile integration.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "10. Assumptions and Open Items", 1)
    for item in [
        "Customer will provide acceptable sample or redacted identity documents for testing.",
        "Formal authoritative-source validation may require government, bank, or trusted data-provider integration.",
        "Production deployment should use HTTPS, a real database, encrypted storage, rate limiting, and monitoring.",
        "Face templates and verification records must follow a written retention and deletion policy.",
        "Manual review should be available for borderline document and biometric cases.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "11. Recommendation", 1)
    add_para(
        doc,
        "Proceed with a controlled pilot for internal users and customer demonstration. Use the prepared NIST IAL2 test matrix to measure accuracy, spoof rejection, document fraud rejection, and usability before expanding to production or regulated payment use cases.",
    )

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
