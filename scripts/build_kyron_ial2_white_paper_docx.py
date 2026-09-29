from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "kyron-ekyc-nist-ial2-white-paper-draft.docx"
LOGO = ROOT / "kyron.png"

NAVY = RGBColor(4, 20, 58)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(94, 108, 132)
GOLD = RGBColor(224, 183, 38)
GREEN = RGBColor(20, 126, 90)
RED = RGBColor(154, 28, 28)
LIGHT_BLUE = "E8EEF7"
LIGHT_GOLD = "FFF8E2"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = 120


def set_run_font(run, *, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, *, before=0, after=6, line=1.10):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_para(doc, text="", *, size=11, color=None, bold=False, italic=False, after=6, before=0, align=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color or NAVY, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, after=4, line=1.167)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(p, after=4, line=1.167)
    p.add_run(text)
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=CELL_MARGIN_DXA, end=CELL_MARGIN_DXA):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, *, indent_dxa=TABLE_INDENT_DXA):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_text(cell, text, *, bold=False, color=NAVY, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.05)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold)


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    for idx, header in enumerate(headers):
        shade_cell(table.rows[0].cells[idx], LIGHT_GRAY)
        set_cell_text(table.rows[0].cells[idx], header, bold=True, size=9.5)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            set_cell_text(row.cells[idx], value, size=9)
    add_para(doc, "", after=4)
    return table


def add_callout(doc, title, body, *, fill=LIGHT_GOLD, accent=GOLD):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.rows[0].cells[0]
    shade_cell(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=2, line=1.08)
    r = p.add_run(title)
    set_run_font(r, size=10.5, color=NAVY, bold=True)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0, line=1.10)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10, color=MUTED)
    add_para(doc, "", after=4)
    return table


def paragraph_bottom_rule(paragraph, color="E0B726", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = NAVY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, RGBColor(31, 77, 120), 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(header_p, after=0)
    run = header_p.add_run("Kyron eKYC White Paper Draft")
    set_run_font(run, size=9, color=MUTED, bold=True)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(footer_p, after=0)
    r = footer_p.add_run("Draft for internal review | Not a certification claim")
    set_run_font(r, size=8.5, color=MUTED)


def add_cover(doc):
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(1.35))
        set_paragraph_spacing(p, after=18)

    kicker = add_para(
        doc,
        "WHITE PAPER DRAFT",
        size=11,
        color=GOLD,
        bold=True,
        after=6,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    title = add_para(
        doc,
        "Kyron eKYC",
        size=30,
        color=NAVY,
        bold=True,
        after=3,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    subtitle = add_para(
        doc,
        "NIST IAL2-Aligned Remote Identity Proofing with Document Fraud Detection, Liveness, and Face ID Authentication",
        size=14,
        color=MUTED,
        after=14,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    paragraph_bottom_rule(subtitle, color="E0B726", size="8")

    add_para(
        doc,
        f"Prepared for internal product, security, and customer review | {date.today().isoformat()}",
        size=10.5,
        color=MUTED,
        after=20,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_callout(
        doc,
        "Positioning statement",
        "Kyron should be described as NIST IAL2-aligned or designed toward NIST IAL2 readiness. "
        "This draft does not assert NIST certification, ISO certification, or external laboratory validation.",
        fill=LIGHT_BLUE,
    )

    add_heading(doc, "Executive Summary", 1)
    add_para(
        doc,
        "Kyron eKYC is a web-first identity verification platform designed to verify document evidence, "
        "bind the applicant to the document portrait through face matching, test live presence through active and "
        "passive liveness controls, and support returning-user Face ID authentication for future banking or payment-style flows.",
    )
    add_para(
        doc,
        "The product direction is aligned with the structure of NIST SP 800-63-4: identity proofing is addressed under "
        "Identity Assurance Level 2 (IAL2), returning-user authentication is addressed separately under Authentication "
        "Assurance Level controls, and biometric attack resistance is evaluated through presentation attack detection testing.",
    )
    add_para(
        doc,
        "The current implementation is a strong prototype, but a formal production claim requires additional evidence: "
        "privacy notices and consent records, encrypted biometric-template storage, authenticated administrative APIs, "
        "audit logs, calibrated model thresholds, independent PAD and biometric testing, and security operations controls.",
    )
    doc.add_page_break()


def add_body(doc):
    add_heading(doc, "1. Scope and Intended Use", 1)
    add_para(
        doc,
        "This white paper frames Kyron as a remote onboarding and returning-user authentication system for organizations "
        "that need document proofing, live-user assurance, and biometric profile reuse. The current workflow is suitable "
        "for demonstrations and controlled pilots. Production financial services use should add multi-factor authentication, "
        "device binding, transaction confirmation, and formal security/privacy governance.",
    )
    add_bullet(doc, "Primary onboarding evidence: passport and Lao ID card capture from upload or camera.")
    add_bullet(doc, "Primary biometric checks: active face liveness, passive selfie liveness, face comparison, and Face ID login.")
    add_bullet(doc, "Primary decision output: pass, reject, or pending with explainable scores and reason codes.")
    add_bullet(doc, "Storage principle: store verification results and verified profile fields; avoid raw media retention by default.")

    add_heading(doc, "2. Standards Basis", 1)
    add_para(
        doc,
        "The standards below provide the recommended reference model for the product white paper and future compliance roadmap.",
    )
    add_table(
        doc,
        ["Standard / Guidance", "Relevance to Kyron", "How it should be used"],
        [
            (
                "NIST SP 800-63-4",
                "Digital identity framework covering identity proofing, authentication, and federation.",
                "Use as the top-level framework. Avoid saying certified unless independently assessed.",
            ),
            (
                "NIST SP 800-63A",
                "Enrollment and identity proofing, including IAL2 proofing expectations.",
                "Map document evidence, validation, biometric binding, fraud checks, privacy, and redress.",
            ),
            (
                "NIST SP 800-63B",
                "Authentication and authenticator management.",
                "Use for returning Face ID login, MFA, replay resistance, throttling, and authenticator binding.",
            ),
            (
                "ISO/IEC 30107-3",
                "Biometric presentation attack detection testing and reporting.",
                "Use to evaluate phone screen, tablet, printed photo, mask, and video replay resistance.",
            ),
            (
                "ISO/IEC 19795-1",
                "Biometric performance testing and reporting.",
                "Use for face match false match, false non-match, and demographic/device performance reports.",
            ),
            (
                "ISO/IEC 24745",
                "Biometric information protection.",
                "Use for encrypted, revocable, and privacy-protective face-template storage.",
            ),
            (
                "ISO/IEC 27001 and 27701",
                "Information security and privacy management systems.",
                "Use for governance, policies, risk management, incident response, and privacy controls.",
            ),
            (
                "FATF Digital Identity Guidance",
                "Risk-based digital identity guidance for customer due diligence.",
                "Use for financial-sector positioning, risk scoring, and ongoing due diligence context.",
            ),
        ],
        [2100, 3000, 4260],
    )

    add_heading(doc, "3. Target End-to-End Workflow", 1)
    add_number(doc, "Start: user chooses new-user signup or returning-user Face ID login.")
    add_number(doc, "Document capture: user uploads or captures passport / Lao ID evidence with camera guidance.")
    add_number(doc, "Document analysis: backend checks image quality, document likeness, OCR/MRZ fields, expiry, and fraud signals.")
    add_number(doc, "Active liveness: user performs randomized face actions such as blink, turn left/right, or open mouth.")
    add_number(doc, "Gesture challenge: user performs randomized hand gestures in a randomized on-screen target area.")
    add_number(doc, "Selfie verification: backend checks selfie quality, passive liveness, screen/print replay cues, and face match.")
    add_number(doc, "Risk decision: app returns pass, reject, or pending with reason codes and scores.")
    add_number(doc, "Enrollment: if passed, user can enroll one active Face ID linked to one verified profile.")
    add_number(doc, "Returning login or payment flow: user submits a live face check and matches against the stored profile.")

    add_heading(doc, "4. NIST IAL2 Alignment Matrix", 1)
    add_table(
        doc,
        ["IAL2 area", "Current Kyron capability", "Required hardening for stronger readiness"],
        [
            (
                "Identity evidence capture",
                "Camera/upload capture for passport and Lao ID; upload size and content-type controls.",
                "Add capture-quality guidance, device metadata, secure capture telemetry, and clear retention notices.",
            ),
            (
                "Evidence validation",
                "OCR/MRZ extraction, expiry checks, document-shape and fraud heuristics.",
                "Add authoritative or cryptographic validation where available, stronger security-feature checks, and validation logs.",
            ),
            (
                "Applicant binding",
                "Document portrait embedding compared with live selfie embedding.",
                "Calibrate face-match thresholds with formal FMR/FNMR reports and demographic/device testing.",
            ),
            (
                "Liveness and PAD",
                "Active face challenges, gesture challenge, passive anti-spoofing, and multi-frame selfie burst.",
                "Run ISO/IEC 30107-3-style PAD evaluation on print, screen, tablet, video replay, and mask attacks.",
            ),
            (
                "Fraud decisioning",
                "Automatic pass/reject/pending, scores, and reason codes.",
                "Add manual review, redress workflow, model-version evidence, and threshold-change approval records.",
            ),
            (
                "Privacy and consent",
                "No raw media storage by default; public demo warning.",
                "Add explicit consent capture, biometric-use notice, retention limits, deletion/export workflows, and DPIA.",
            ),
            (
                "Records and auditability",
                "Verification result metadata and profile records.",
                "Add immutable audit logs for enrollment, decisions, login, deletion, admin access, model versions, and risk changes.",
            ),
        ],
        [1900, 3100, 4360],
    )

    add_heading(doc, "5. Security and Fraud Detection Controls", 1)
    add_heading(doc, "5.1 Document Fraud Detection", 2)
    add_para(
        doc,
        "The document layer should detect expired documents, unreadable MRZ/OCR, non-document uploads, printed copies, "
        "screen recapture, portrait substitution, text overlay, image manipulation, and low-quality captures. For IAL2 readiness, "
        "Kyron should separate weak-quality failures from high-confidence fraud failures so genuine users are not over-rejected.",
    )
    add_heading(doc, "5.2 Active Liveness", 2)
    add_para(
        doc,
        "Active liveness should require randomized actions and server-verified evidence bursts. The backend should not trust "
        "frontend landmark motion alone. Strong replay signals, repeated static display texture, or visible screen/tablet cues "
        "should cause rejection or manual review.",
    )
    add_heading(doc, "5.3 Passive Liveness and Selfie Burst", 2)
    add_para(
        doc,
        "Passive selfie liveness should use a short burst rather than a single frame. This lets the system compare texture, "
        "motion, lighting variation, repeated display edges, glare, moire patterns, and model spoof scores across frames.",
    )
    add_heading(doc, "5.4 Face Matching and Face ID", 2)
    add_para(
        doc,
        "Face matching should report calibrated genuine-match and impostor-match performance. Returning-user Face ID should be "
        "treated as a biometric factor, not as a complete bank-grade authenticator by itself. For payment or transfer, Kyron should "
        "combine Face ID with device binding, PIN/passkey/WebAuthn, transaction confirmation, rate limits, and risk scoring.",
    )

    add_heading(doc, "6. Biometric Data Protection", 1)
    add_callout(
        doc,
        "Production requirement",
        "Face templates and verified identity attributes should be encrypted at rest, protected by access controls, covered by "
        "explicit consent, and governed by retention/deletion workflows. Raw identity media should remain disabled by default.",
        fill=LIGHT_BLUE,
    )
    add_bullet(doc, "Encrypt face templates and sensitive profile attributes using KMS or envelope encryption.")
    add_bullet(doc, "Record consent version, timestamp, biometric purpose, and retention period.")
    add_bullet(doc, "Support account deletion, biometric-template deletion, profile export, and retention expiry.")
    add_bullet(doc, "Protect admin endpoints with authentication, authorization, throttling, and audit logging.")
    add_bullet(doc, "Separate demo/test data from production data and prevent public demo uploads of sensitive real documents.")

    add_heading(doc, "7. Testing and Evidence Plan", 1)
    add_para(
        doc,
        "A credible IAL2-alignment claim depends on evidence. The project should maintain test datasets, evaluation scripts, "
        "threshold history, model versions, and repeatable reports.",
    )
    add_table(
        doc,
        ["Test domain", "Representative cases", "Metrics to report"],
        [
            (
                "Document proofing",
                "Real camera passport, real camera Lao ID, printed passport copy, printed ID copy, edited portrait, text overlay, expired evidence, non-document upload.",
                "Accuracy, false reject rate, false accept rate, MRZ/OCR read rate, expiry extraction accuracy.",
            ),
            (
                "Active liveness",
                "Real face, phone replay, tablet/laptop replay, video replay, low light, glasses, partial face, wrong action.",
                "Attack presentation classification error rate, bona fide rejection, challenge pass latency.",
            ),
            (
                "Passive selfie PAD",
                "Real live selfie, phone screen photo, tablet/laptop screen, printed photo, paper attack, video replay, no-valid-face.",
                "PAD pass/fail accuracy, false accept rate, false reject rate, score distributions by device.",
            ),
            (
                "Face matching",
                "Same-person document/selfie pairs, impostor pairs, look-alike pairs, aging/lighting/device variation.",
                "False match rate, false non-match rate, threshold curve, demographic/device breakdown.",
            ),
        ],
        [1750, 3900, 3710],
    )
    add_para(
        doc,
        "Face login should be tested separately with returning users, wrong users, duplicate-profile attempts, spoof login attempts, "
        "multiple-face frames, low-confidence faces, lockout behavior, and average login latency.",
        size=10,
        color=MUTED,
        after=6,
    )

    add_heading(doc, "8. Production Readiness Roadmap", 1)
    add_heading(doc, "Phase 1: Security Hardening", 2)
    add_bullet(doc, "Protect profile/admin endpoints with authenticated role-based access.")
    add_bullet(doc, "Encrypt biometric templates and sensitive identity attributes.")
    add_bullet(doc, "Add audit log tables and event capture for decisions, enrollment, login, deletion, and admin actions.")
    add_bullet(doc, "Add consent, retention, deletion, and privacy-notice workflows.")
    add_heading(doc, "Phase 2: Model Calibration and Dataset Evidence", 2)
    add_bullet(doc, "Finalize document, selfie PAD, active liveness, and face-match test datasets.")
    add_bullet(doc, "Generate repeatable accuracy reports and threshold recommendations.")
    add_bullet(doc, "Document false accepts, false rejects, known edge cases, and manual-review triggers.")
    add_heading(doc, "Phase 3: Operational Controls", 2)
    add_bullet(doc, "Add monitoring, alerting, backups, incident response, secrets management, and dependency scanning.")
    add_bullet(doc, "Run penetration testing and privacy impact assessment before broader customer use.")
    add_bullet(doc, "Prepare ISO/IEC 27001/27701 policies if the organization wants formal certification.")
    add_heading(doc, "Phase 4: Independent Review", 2)
    add_bullet(doc, "Perform external PAD/biometric evaluation against ISO/IEC 30107-3 and ISO/IEC 19795-1 style reporting.")
    add_bullet(doc, "Review NIST IAL2 alignment matrix with an identity/security assessor.")
    add_bullet(doc, "Update customer-facing claims only after evidence is collected and reviewed.")

    add_heading(doc, "9. Recommended Customer-Facing Claim Language", 1)
    add_callout(
        doc,
        "Recommended wording",
        "Kyron eKYC is designed as a NIST IAL2-aligned remote identity proofing workflow using document evidence validation, "
        "active/passive liveness, face matching, and explainable fraud-risk decisioning. Certification and conformity claims "
        "should be made only after independent assessment.",
        fill=LIGHT_GOLD,
    )
    add_para(
        doc,
        "Avoid wording: do not say NIST IAL2 certified, ISO certified, bank-grade certified, or government-approved unless an external assessment is complete.",
        size=9.5,
        color=RED,
        bold=True,
        after=8,
    )

    add_heading(doc, "10. Current Limitations", 1)
    add_bullet(doc, "Document validation is still model- and heuristic-driven; authoritative issuer validation is not yet integrated.")
    add_bullet(doc, "Face templates require production-grade encryption and key management.")
    add_bullet(doc, "Profile admin APIs are suitable for demos only until protected by authentication and authorization.")
    add_bullet(doc, "PAD and face-match thresholds need larger datasets and documented false reject / false accept tradeoffs.")
    add_bullet(doc, "Payment or transfer flows require MFA and transaction controls; face alone is not enough for high-risk banking use.")
    add_bullet(doc, "Formal claims require independent review, audit evidence, privacy documentation, and operational security controls.")

    add_heading(doc, "References", 1)
    refs = [
        "NIST SP 800-63-4, Digital Identity Guidelines: https://pages.nist.gov/800-63-4/sp800-63.html",
        "NIST SP 800-63A, Enrollment and Identity Proofing: https://pages.nist.gov/800-63-4/sp800-63a.html",
        "NIST SP 800-63B, Authentication and Authenticator Management: https://pages.nist.gov/800-63-4/sp800-63b.html",
        "NIST IAL2 Remote Identity Proofing implementation resources: https://pages.nist.gov/800-63-3-Implementation-Resources/63A/ial2remote/",
        "ISO/IEC 30107-3:2023, Biometric presentation attack detection: https://www.iso.org/standard/79520.html",
        "ISO/IEC 19795-1:2021, Biometric performance testing and reporting: https://www.iso.org/standard/73515.html",
        "ISO/IEC 24745:2022, Biometric information protection: https://www.iso.org/standard/75302.html",
        "ISO/IEC 27001:2022, Information security management systems: https://www.iso.org/standard/27001",
        "ISO/IEC 27701:2025, Privacy information management systems: https://www.iso.org/standard/85819.html",
        "FATF Guidance on Digital Identity: https://www.fatf-gafi.org/en/publications/Financialinclusionandnpoissues/Digital-identity-guidance.html",
    ]
    for ref in refs:
        add_bullet(doc, ref)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_body(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
